"""报告 DOCX 渲染器离线单测（C18）。"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

from docx import Document

from ai_engine.agents.auditor import (
    EnhancedRiskItem,
    PreciseCitation,
    RiskCategory,
    RiskLevel,
)
from ai_engine.agents.conflict import ConflictType, CrossClauseConflict
from worker.report import ReportRenderInput, render_report_docx

_WORKER_ROOT = Path(__file__).resolve().parents[1]


def _citation(source_title: str, content: str) -> PreciseCitation:
    return PreciseCitation(
        sourceType="STANDARD_CLAUSE",
        sourceTitle=source_title,
        contentChunk=content,
        relevanceScore=0.93,
    )


def _risk(title: str = "付款期限风险") -> EnhancedRiskItem:
    return EnhancedRiskItem(
        id="risk-1",
        clauseId="clause-1",
        category=RiskCategory.PAYMENT_TERM_RISK,
        level=RiskLevel.HIGH,
        title=title,
        originalText="甲方应在 30 日内付款。",
        bbox={},
        suggestion="明确付款节点与逾期责任。",
        legalBasis="《民法典》相关条款",
        citations=[
            _citation("标准付款条款", "付款期限应当明确、无歧义。"),
            _citation("付款期限判例", "合同应约定具体付款节点。"),
        ],
    )


def _conflict() -> CrossClauseConflict:
    return CrossClauseConflict(
        id="conflict-1",
        sourceClauseId="clause-2",
        targetClauseId="clause-3",
        conflictType=ConflictType.TERM_CONTRADICTION,
        description="两条款对交付期限约定不一致。",
        level=RiskLevel.MEDIUM,
        suggestion="统一交付期限表述。",
        citations=[_citation("交付期限标准条款", "交付期限只能有一个生效口径。")],
    )


def _input(
    *,
    risks: list[EnhancedRiskItem] | None = None,
    conflicts: list[CrossClauseConflict] | None = None,
) -> ReportRenderInput:
    return ReportRenderInput(
        contract_id="contract-18",
        generated_at=datetime(2026, 9, 27, 10, 0, 0, tzinfo=UTC),
        risks=risks if risks is not None else [_risk()],
        conflicts=conflicts if conflicts is not None else [_conflict()],
        elements=[
            {
                "key": "partyA",
                "label": "甲方",
                "value": "示例公司",
                "confidence": 0.99,
            }
        ],
        health_score={
            "overall": 82,
            "dimensions": {
                "compliance": 85,
                "performanceRisk": 78,
                "clauseCompleteness": 83,
            },
        },
    )


def _paragraph_texts(data: bytes) -> list[str]:
    document = Document(BytesIO(data))
    return [paragraph.text for paragraph in document.paragraphs]


def test_render_input_accepts_payload_contract_aliases() -> None:
    parsed = ReportRenderInput.model_validate(
        {
            "contractId": "contract-18-alias",
            "generatedAt": "2026-09-27T10:00:00+00:00",
            "risks": [],
            "conflicts": [],
            "elements": [],
            "healthScore": {"overall": 80, "dimensions": {}},
        }
    )

    assert parsed.contract_id == "contract-18-alias"
    assert parsed.health_score == {"overall": 80, "dimensions": {}}


def test_render_returns_openable_docx() -> None:
    data = render_report_docx(_input())

    assert isinstance(data, bytes)
    assert len(data) > 0
    assert _paragraph_texts(data)


def test_render_contains_risk_and_citations() -> None:
    texts = _paragraph_texts(render_report_docx(_input()))
    joined = "\n".join(texts)

    assert "付款期限风险" in joined
    assert "标准付款条款" in joined
    assert "付款期限应当明确、无歧义。" in joined
    assert "付款期限判例" in joined
    assert "合同应约定具体付款节点。" in joined


def test_render_contains_conflict_and_citation() -> None:
    texts = _paragraph_texts(render_report_docx(_input()))
    joined = "\n".join(texts)

    assert "矛盾 1" in joined
    assert "交付期限标准条款" in joined
    assert "交付期限只能有一个生效口径。" in joined


def test_render_counts_match_input() -> None:
    data = render_report_docx(
        _input(
            risks=[_risk("风险A"), _risk("风险B")],
            conflicts=[_conflict(), _conflict()],
        )
    )
    texts = _paragraph_texts(data)

    assert sum(text.startswith("风险 ") for text in texts) == 2
    assert sum(text.startswith("矛盾 ") for text in texts) == 2
    assert any(text.startswith("风险 1：风险A") for text in texts)
    assert any(text.startswith("风险 2：风险B") for text in texts)


def test_render_empty_sections_are_still_valid() -> None:
    data = render_report_docx(
        _input(
            risks=[],
            conflicts=[],
        )
    )
    texts = _paragraph_texts(data)
    joined = "\n".join(texts)

    assert "风险清单" in joined
    assert "矛盾清单" in joined
    assert "无" in texts
    assert Document(BytesIO(data)).paragraphs


def test_render_contains_health_score() -> None:
    texts = _paragraph_texts(render_report_docx(_input()))
    joined = "\n".join(texts)

    assert "健康分" in joined
    assert "总分：82" in joined
    assert "compliance：85" in joined


def test_renderer_has_no_forbidden_imports() -> None:
    source = (_WORKER_ROOT / "worker" / "report" / "render_docx.py").read_text(
        encoding="utf-8"
    )

    assert "from app." not in source
    assert "import app." not in source
    assert "import minio" not in source
    assert "import sqlalchemy" not in source
