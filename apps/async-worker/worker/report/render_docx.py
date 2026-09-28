"""审核报告 DOCX 渲染器。

设计要点：
1. 只消费 `ai_engine` 领域模型，不 import 网关 DTO、数据库或对象存储 SDK。
2. 使用 `python-docx` 对象 API，禁止拼接 OOXML 字符串。
3. 空输入也返回合法 DOCX，保证报告生成任务始终有可交付产物。

对应：`文件集合/codex/4.C18报告DOCX渲染器.md`。
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from pydantic import BaseModel, ConfigDict, Field

from ai_engine.agents.auditor import EnhancedRiskItem
from ai_engine.agents.conflict import CrossClauseConflict


class ReportRenderInput(BaseModel):
    """报告渲染输入（与 `4-3 §⑤.1` 载荷契约对齐）。"""

    model_config = ConfigDict(populate_by_name=True)

    contract_id: str = Field(alias="contractId")
    generated_at: datetime = Field(alias="generatedAt")
    risks: list[EnhancedRiskItem] = Field(default_factory=list)
    conflicts: list[CrossClauseConflict] = Field(default_factory=list)
    elements: list[dict[str, Any]] = Field(default_factory=list)
    health_score: dict[str, Any] | None = Field(default=None, alias="healthScore")


def render_report_docx(payload: ReportRenderInput) -> bytes:
    """渲染审核报告并返回 DOCX 字节。"""
    document = Document()
    _add_header(document, payload)
    _add_risks(document, payload.risks)
    _add_conflicts(document, payload.conflicts)
    _add_elements(document, payload.elements)
    _add_health_score(document, payload.health_score)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_header(document: DocumentObject, payload: ReportRenderInput) -> None:
    """写入标题、合同编号与生成时间。"""
    document.add_heading("合同智能审核报告", level=1)
    _add_field(document, "合同编号", payload.contract_id)
    _add_field(document, "生成时间", _format_datetime(payload.generated_at))


def _add_risks(
    document: DocumentObject,
    risks: list[EnhancedRiskItem],
) -> None:
    """逐条渲染风险项及其溯源。"""
    document.add_heading("风险清单", level=1)
    if not risks:
        document.add_paragraph("无")
        return

    for index, risk in enumerate(risks, start=1):
        document.add_heading(f"风险 {index}：{risk.title}", level=2)
        _add_field(document, "等级", risk.level.value)
        _add_field(document, "类别", risk.category.value)
        _add_field(document, "条款", risk.clauseId)
        _add_field(document, "原文", risk.originalText)
        _add_field(document, "建议", risk.suggestion)
        _add_field(document, "法律依据", risk.legalBasis)
        _add_citations(document, risk.citations)


def _add_conflicts(
    document: DocumentObject,
    conflicts: list[CrossClauseConflict],
) -> None:
    """逐条渲染跨条款矛盾及其溯源。"""
    document.add_heading("矛盾清单", level=1)
    if not conflicts:
        document.add_paragraph("无")
        return

    for index, conflict in enumerate(conflicts, start=1):
        document.add_heading(f"矛盾 {index}", level=2)
        _add_field(document, "来源条款", conflict.sourceClauseId)
        _add_field(document, "目标条款", conflict.targetClauseId)
        _add_field(document, "矛盾类型", conflict.conflictType.value)
        _add_field(document, "等级", conflict.level.value)
        _add_field(document, "说明", conflict.description)
        _add_field(document, "建议", conflict.suggestion)
        _add_citations(document, conflict.citations)


def _add_elements(
    document: DocumentObject,
    elements: list[dict[str, Any]],
) -> None:
    """写入合同要素摘要。"""
    document.add_heading("合同要素", level=1)
    if not elements:
        document.add_paragraph("无")
        return

    for index, element in enumerate(elements, start=1):
        label = element.get("label") or element.get("key") or f"要素 {index}"
        value = element.get("value")
        confidence = element.get("confidence")
        text = f"{label}：{value if value is not None else '未抽取'}"
        if confidence is not None:
            text += f"（置信度 {float(confidence):.2f}）"
        document.add_paragraph(text)


def _add_health_score(
    document: DocumentObject,
    health_score: dict[str, Any] | None,
) -> None:
    """写入可选的健康分区块。"""
    if not health_score:
        return

    document.add_heading("健康分", level=1)
    _add_field(document, "总分", health_score.get("overall", "未提供"))
    dimensions = health_score.get("dimensions")
    if not isinstance(dimensions, dict):
        return

    for name in ("compliance", "performanceRisk", "clauseCompleteness"):
        _add_field(document, name, dimensions.get(name, "未提供"))


def _add_citations(
    document: DocumentObject,
    citations: list[Any],
) -> None:
    """逐条渲染精确溯源。"""
    document.add_heading("溯源", level=3)
    if not citations:
        document.add_paragraph("无")
        return

    for index, citation in enumerate(citations, start=1):
        document.add_paragraph(
            f"[{index}] {citation.sourceTitle}：{citation.contentChunk}"
            f"（相关度 {float(citation.relevanceScore):.2f}）",
            style="List Bullet",
        )


def _add_field(
    document: DocumentObject,
    label: str,
    value: Any,
) -> None:
    """写入一个可读字段，空值统一显示为“未提供”。"""
    display_value = value if value is not None else "未提供"
    document.add_paragraph(f"{label}：{display_value}")


def _format_datetime(value: datetime) -> str:
    """把生成时间格式化为稳定可读文本。"""
    if value.tzinfo is not None:
        return value.strftime("%Y-%m-%d %H:%M:%S %Z")
    return value.strftime("%Y-%m-%d %H:%M:%S")


__all__ = ["ReportRenderInput", "render_report_docx"]
